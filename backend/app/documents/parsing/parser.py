from io import BytesIO
from pathlib import Path
import yaml
from docx import Document
from pypdf import PdfReader

MAX_BYTES = 5 * 1024 * 1024
MAX_CHARS = 120000


def parse(data: bytes, filename: str):
    if not data or len(data) > MAX_BYTES:
        raise ValueError("File must contain between 1 byte and 5 MiB")
    suffix = Path(filename).suffix.lower()
    if suffix in {".md", ".txt"}:
        text = data.decode("utf-8")
    elif suffix == ".pdf":
        reader = PdfReader(BytesIO(data))
        if len(reader.pages) > 200:
            raise ValueError("PDF exceeds 200 pages")
        text = "\n\n".join(p.extract_text() or "" for p in reader.pages)
    elif suffix == ".docx":
        import zipfile

        with zipfile.ZipFile(BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > 20 * 1024 * 1024:
                raise ValueError("DOCX expanded size exceeds 20 MiB")
        doc = Document(BytesIO(data))
        text = "\n\n".join(
            [p.text for p in doc.paragraphs]
            + [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
        )
    else:
        raise ValueError("Supported formats: Markdown, text, PDF, DOCX")
    if not text.strip() or len(text) > MAX_CHARS:
        raise ValueError("Document must contain extractable text, at most 120,000 characters")
    metadata = None
    if text.startswith("---\n") or text.startswith("---\r\n"):
        parts = text.replace("\r\n", "\n").split("---\n", 2)
        if len(parts) != 3:
            raise ValueError("Invalid document front matter")
        metadata = yaml.safe_load(parts[1])
        text = parts[2]
    return metadata, text.strip()
